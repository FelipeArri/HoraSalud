"""Pruebas del padrón FONASA simulado (base externa separada) y de los comandos de consulta."""
import tempfile
from io import StringIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import connections
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from autenticacionApp import inscripcion
from autenticacionApp.models import Paciente, Usuario

from .models import PersonaFonasa


class PadronTests(TestCase):
    databases = {'default', 'fonasa'}

    def test_padron_inicial_con_25_personas_en_su_propia_base(self):
        self.assertEqual(PersonaFonasa.objects.count(), 25)
        tablas_horasalud = connections['default'].introspection.table_names()
        self.assertNotIn(PersonaFonasa._meta.db_table, tablas_horasalud)  # sistemas separados

    def test_persona_agregada_en_el_padron_puede_recibir_codigo(self):
        PersonaFonasa.objects.create(rut='31000020-5', nombres='Nueva', apellido_paterno='Agregada',
                                     prevision='FONASA', tramo='C', sector='SECTOR 3')
        consulta = inscripcion.consultar('31000020-5')
        self.assertTrue(consulta.puede_recibir_codigo)
        self.assertEqual((consulta.inscripcion['sector'], consulta.tramo_texto), ('Sector 3', 'FONASA tramo C'))

    def test_inscrito_en_otro_cesfam_no_corresponde(self):
        self.assertTrue(PersonaFonasa.objects.filter(rut='31000013-2', establecimiento='CESFAM Santa Anselma').exists())
        self.assertIn('Debe inscribirse', inscripcion.consultar('31000013-2').motivo_rechazo)

    def test_validaciones_del_padron(self):
        with self.assertRaises(ValidationError):
            PersonaFonasa(rut='31000020-9', nombres='X', apellido_paterno='Y', tramo='A').full_clean()  # DV malo
        with self.assertRaises(ValidationError):
            PersonaFonasa(rut='31000021-3', nombres='X', apellido_paterno='Y', prevision='FONASA').full_clean()  # sin tramo


class ComandosExcelTests(TestCase):
    databases = {'default', 'fonasa'}

    def setUp(self):
        call_command('cargar_datos_demo', stdout=StringIO())
        self.carpeta = Path(tempfile.mkdtemp())

    def test_exportar_padron(self):
        with override_settings(REPORTES_DIR=self.carpeta):
            call_command('exportar_padron', stdout=StringIO())
        hoja = load_workbook(self.carpeta / 'padron_fonasa_simulado.xlsx').active
        ruts = [hoja.cell(f, 1).value for f in range(6, hoja.max_row + 1)]
        self.assertEqual(len(ruts), 25)
        self.assertIn('31000002-7', ruts)

    def test_exportar_inscritos(self):
        with override_settings(REPORTES_DIR=self.carpeta):
            call_command('exportar_inscritos', stdout=StringIO())
        archivo = next(self.carpeta.glob('inscritos_*.xlsx'))
        hoja = load_workbook(archivo).active
        self.assertEqual(hoja.max_row - 5, Usuario.objects.filter(tipo_usuario__in=['paciente', 'tutor']).count())

    def test_vaciar_y_empezar_sin_pacientes(self):
        call_command('vaciar_pacientes', '--confirmar', stdout=StringIO())
        self.assertFalse(Paciente.objects.exists())
        self.assertFalse(Usuario.objects.filter(tipo_usuario__in=['paciente', 'tutor']).exists())
        self.assertTrue(Usuario.objects.filter(tipo_usuario='profesional').exists())  # funcionarios intactos
        self.assertEqual(PersonaFonasa.objects.count(), 25)  # padrón intacto
        call_command('cargar_datos_demo', '--sin-pacientes', stdout=StringIO())
        self.assertFalse(Paciente.objects.exists())
