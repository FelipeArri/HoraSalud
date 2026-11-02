"""Utilidad para generar los Excel de consulta (reportes/), con el estilo del proyecto."""
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERDE = PatternFill('solid', fgColor='419800')
BORDE = Border(*(Side(style='thin', color='999999'),) * 4)


def guardar_excel(nombre_archivo, titulo, subtitulo, encabezados, filas, anchos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja['A1'] = 'PROYECTO HORASALUD — INGENIERÍA DE SOFTWARE'
    hoja['A1'].font = Font(bold=True, size=12)
    hoja['A2'] = titulo
    hoja['A2'].font = Font(bold=True, size=14, color='1F4E79')
    hoja['A3'] = subtitulo
    hoja['A3'].font = Font(italic=True, color='555555')
    fila_enc = 5
    for col, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(fila_enc, col, texto)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = VERDE
        celda.border = BORDE
        celda.alignment = Alignment(vertical='center', wrap_text=True)
    for i, fila in enumerate(filas, start=fila_enc + 1):
        for col, valor in enumerate(fila, start=1):
            celda = hoja.cell(i, col, valor)
            celda.border = BORDE
            celda.alignment = Alignment(vertical='top', wrap_text=True)
    for col, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.freeze_panes = hoja.cell(fila_enc + 1, 1)
    hoja.auto_filter.ref = f'A{fila_enc}:{get_column_letter(len(encabezados))}{fila_enc + len(filas)}'
    carpeta = Path(settings.REPORTES_DIR)
    carpeta.mkdir(exist_ok=True)
    destino = carpeta / nombre_archivo
    libro.save(destino)
    return destino
